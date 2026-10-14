"""Builds the Storyline-import PowerPoint for AI Literacy Level 1 on the Resource Pro template.

Usage: python3 build-deck.py <PPt_template.pptx> <deck-data.json> <out.pptx>

Every screen becomes one slide with editable, named shapes; voiceover goes in the speaker notes.
Layers (Meera, feedback, tabs, results) become extra slides marked LAYER, to copy into Storyline layers.
"""
import json, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

TPL, DATA, OUT = sys.argv[1:4]
data = json.load(open(DATA))
SCREENS = data['screens']

C = {k: RGBColor.from_string(v) for k, v in dict(
    navy='021F62', teal='00BBC2', green='7EBD31', blue='497EDC', bluesoft='E8EEFB', tealsoft='D9F5F6',
    tealink='00787D', greentxt='4F7A1C', greensoft='EEF5E1', coral='FB4671', coralink='C42A50', coralsoft='FFE6EC',
    grey='5A5C61', line='D8DBE2', snow='F6F6F9', tile='0F2D74', tileline='2B3F73', lightblue='9DB8EE', white='FFFFFF',
    orange='E07A1F').items()}
HEAD, BODY, MONO = 'Aptos Display', 'Aptos', 'Aptos Mono'
px = lambda v: Emu(int(round(v * 9525)))          # 1280 x 720 px canvas = 13.333 x 7.5 in
import os
ICONS = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'assets', 'icons')   # rendered from the prototype by render-icons.cjs

prs = Presentation(TPL)
lst = prs.slides._sldIdLst
for sid in list(lst):                              # keep masters and layouts, drop the template's sample slides
    prs.part.drop_rel(sid.rId); lst.remove(sid)
L_COVER = prs.slide_masters[0].slide_layouts[0]
L_LIGHT = prs.slide_masters[0].slide_layouts[5]    # BASIC SLIDE
L_DARK = prs.slide_masters[1].slide_layouts[2]     # EMPTY SLIDE (navy)

# ---------- drawing helpers ----------
def _alpha(shape, pct_opaque):
    sf = shape.fill._xPr.find(qn('a:solidFill'))
    clr = sf[0]
    a = etree.SubElement(clr, qn('a:alpha')); a.set('val', str(int(pct_opaque * 1000)))

def box(s, x, y, w, h, fill=None, line=None, dash=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, lw=1.5, name=None, alpha=None):
    b = s.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    if name: b.name = name
    if fill is None: b.fill.background()
    else:
        b.fill.solid(); b.fill.fore_color.rgb = C[fill]
        if alpha is not None: _alpha(b, alpha)
    if line is None: b.line.fill.background()
    else:
        b.line.color.rgb = C[line]; b.line.width = Pt(lw)
        if dash:
            ln = b.line._get_or_add_ln(); d = etree.SubElement(ln, qn('a:prstDash')); d.set('val', 'dash')
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        b.adjustments[0] = radius
    b.shadow.inherit = False
    tf = b.text_frame
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    return b

def write(shape_or_tf, runs, size=18, color='navy', bold=False, font=BODY, align='l', anchor='t', italic=False, spacing=None):
    tf = shape_or_tf.text_frame if hasattr(shape_or_tf, 'text_frame') else shape_or_tf
    tf.word_wrap = True
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    paras = runs if isinstance(runs, list) and runs and isinstance(runs[0], list) else [runs if isinstance(runs, list) else [(runs, {})]]
    for pi, para in enumerate(paras):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if spacing: p.space_after = Pt(spacing)
        for t, o in para:
            r = p.add_run(); r.text = t
            f = r.font; f.size = Pt(o.get('size', size)); f.bold = o.get('bold', bold); f.italic = o.get('italic', italic)
            f.color.rgb = C[o.get('color', color)]; f.name = o.get('font', font)
    return shape_or_tf

def text(s, x, y, w, h, runs, name=None, **kw):
    tb = s.shapes.add_textbox(px(x), px(y), px(w), px(h))
    if name: tb.name = name
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tb.text_frame, m, 0)
    write(tb, runs, **kw)
    return tb

def line(s, x1, y1, x2, y2, color='line', w=2, name=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    c.line.color.rgb = C[color]; c.line.width = Pt(w)
    if name: c.name = name
    return c

def icon(s, key, x, y, size, name=None):
    pic = s.shapes.add_picture(os.path.join(ICONS, key + '.png'), px(x), px(y), px(size), px(size))
    pic.name = name or f'Icon {key}'
    return pic

def notes(s, *parts):
    s.notes_slide.notes_text_frame.text = '\n\n'.join(p for p in parts if p)

# ---------- chrome ----------
SEC_LABELS = [x['label'] for x in data['sections']]
def progress(s, sec):
    if sec is None: return
    x0, y0, gap = 75, 16, 100
    line(s, x0 + 6, y0 + 6, x0 + 6 + gap * 7, y0 + 6, 'line', 1.5, 'Progress line')
    if sec > 0: line(s, x0 + 6, y0 + 6, x0 + 6 + gap * sec, y0 + 6, 'teal', 1.5, 'Progress line done')
    for k in range(8):
        if k < sec: d = box(s, x0 + k * gap, y0, 12, 12, 'teal', shape=MSO_SHAPE.OVAL, name=f'Dot {k+1}')
        elif k == sec:
            box(s, x0 + k * gap - 5, y0 - 5, 22, 22, 'tealsoft', shape=MSO_SHAPE.OVAL, name='Current dot halo')
            box(s, x0 + k * gap, y0, 12, 12, 'white', 'teal', shape=MSO_SHAPE.OVAL, lw=1.5, name=f'Dot {k+1}')
        else: box(s, x0 + k * gap, y0, 12, 12, 'line', shape=MSO_SHAPE.OVAL, name=f'Dot {k+1}')
    text(s, x0 + sec * gap - 74, y0 + 15, 160, 14, SEC_LABELS[sec], name='Progress label', size=8, color='grey', align='c')

def badge(s, sc):
    c = sc['cue']
    if not c: return
    sq = box(s, 1169, 96, 36, 36, 'navy', radius=0.15, name='AWARE badge letter')
    write(sq, c['L'], size=20, bold=True, color='teal', font=HEAD, align='c', anchor='m')
    text(s, 860, 104, 300, 20, c['w'], name='AWARE badge word', size=11, color='grey', align='r')

def title_light(s, t, eyebrow=None):
    ph = s.placeholders[10]; ph.text_frame.text = ''
    ph.left, ph.top, ph.width, ph.height = px(74), px(70), px(900), px(44)
    write(ph, t, size=30 if len(t) <= 40 else 26, bold=True, color='navy', font=HEAD)
    ph.name = 'Title'
    if eyebrow: text(s, 75, 50, 700, 16, eyebrow, name='Section label', size=11, bold=True, color='tealink')

def title_dark(s, t, y=60, size=40):
    text(s, 75, y, 1000, 60, t, name='Title', size=size, bold=True, color='green', font=HEAD)

def solid_bg(s, dark):
    # The template's light background is white at 0% opacity; Storyline renders that as black, so set it explicitly.
    s.background.fill.solid(); s.background.fill.fore_color.rgb = C['navy' if dark else 'snow']

def new(sc, dark=False, eyebrow=None):
    s = prs.slides.add_slide(L_DARK if dark else L_LIGHT)
    solid_bg(s, dark)
    if not dark: title_light(s, sc['title'], eyebrow or sc.get('section_name'))
    progress(s, sc['section'] if sc['layout'] not in ('opener', 'assessIntro') else None)
    if not dark: badge(s, sc)
    return s

def layer_slide(sc, name, how, dark=False):
    s = prs.slides.add_slide(L_DARK if dark else L_LIGHT)
    solid_bg(s, dark)
    if not dark:
        ph = s.placeholders[10]; ph.text_frame.text = ''
        write(ph, f"{sc['title']}", size=30, bold=True, color='navy', font=HEAD)
    tag = box(s, 75, 4, 760, 26, 'orange', radius=0.2, name='DELETE ME: layer tag')
    write(tag, f"LAYER · Screen {sc['screen']} · {name}: copy the objects below into the “{name}” layer, then delete this slide",
          size=10, bold=True, color='white', align='c', anchor='m')
    notes(s, f"LAYER for screen {sc['screen']}: {name}", how)
    return s

def card(s, x, y, w, h, term, desc, kind='ai', term_size=14, desc_size=17, name=None):
    fill, ln, dash, tcol = {'ai': ('bluesoft', 'blue', True, 'blue'), 'human': ('tealsoft', 'teal', False, 'tealink'),
                            'white': ('white', 'line', False, 'blue'), 'ok': ('greensoft', 'greentxt', False, 'greentxt')}[kind]
    b = box(s, x, y, w, h, fill, ln, dash=dash, name=name or f'Card {term}')
    b.text_frame.margin_left = b.text_frame.margin_right = px(20); b.text_frame.margin_top = px(16)
    paras = [[(term, dict(size=term_size, bold=True, color=tcol, font=MONO))]]
    if desc: paras.append([(desc, dict(size=desc_size, color='navy'))])
    write(b, paras, spacing=6)
    return b

def arrow(s, x, y, color='teal', size=24):
    text(s, x, y, 40, 30, '→', name='Arrow', size=size, bold=True, color=color, align='c')

def lines_of(onscreen):
    return [l.strip() for l in onscreen.split('\n') if l.strip()]

def std_notes(sc, extra=''):
    return f"VOICEOVER\n{sc['vo']}", (f"BUILD IN STORYLINE\n{extra}" if extra else ''), f"Design note: {sc.get('storyline','')}"

# ---------- layouts ----------
def L_title(sc):
    s = prs.slides.add_slide(L_COVER)
    ph = s.placeholders[10]; ph.text_frame.text = ''
    write(ph, [('AI ', dict(color='white')), ('Literacy', dict(color='green'))], size=54, bold=True, font=HEAD); ph.name = 'Title'
    sub = s.placeholders[11]; sub.text_frame.text = ''
    write(sub, [('Understand AI. ', {}), ('Use It Responsibly.', dict(bold=True))], size=20, color='white'); sub.name = 'Subtitle'
    det = s.placeholders[12]; det.text_frame.text = ''
    write(det, 'Level 1', size=12, color='teal'); det.name = 'Level'
    notes(s, *std_notes(sc, 'Fade the title (0.75 s), then the subtitle 0.5 s later.'))

def L_outcomes(sc):
    s = new(sc)
    text(s, 75, 125, 900, 30, sc['intro'], name='Intro', size=20, color='navy')
    for k, t in enumerate(sc['items']):
        y = 175 + k * 92
        line(s, 75, y, 1205, y, 'line', 1, f'Row line {k+1}')
        c = box(s, 75, y + 18, 52, 52, 'teal', shape=MSO_SHAPE.OVAL, name=f'Number {k+1}')
        write(c, str(k + 1), size=20, bold=True, color='navy', font=HEAD, align='c', anchor='m')
        text(s, 150, y + 26, 1040, 40, t, name=f'Outcome {k+1}', size=22, color='navy', anchor='m')
    notes(s, *std_notes(sc, 'Each row fades in on its cue point.'))

def L_journey(sc):
    s = new(sc)
    line(s, 140, 300, 1140, 300, 'teal', 3, 'Journey line')
    for k, t in enumerate(sc['steps']):
        cx = 140 + k * (1000 / 7)
        c = box(s, cx - 35, 265, 70, 70, 'navy' if k == 7 else 'white', 'teal', shape=MSO_SHAPE.OVAL, lw=3, name=f'Stop {k+1}')
        write(c, str(k + 1), size=22, bold=True, color='teal' if k == 7 else 'navy', font=HEAD, align='c', anchor='m')
        text(s, cx - 70, 350, 140, 60, t, name=f'Stop label {k+1}', size=13, color='navy', align='c')
    notes(s, *std_notes(sc, 'Line fades first, then each stop on its cue.'))

def aware_tiles(s, letters, y=210):
    for k, (L, w) in enumerate(letters):
        x = 75 + k * 230
        t = box(s, x, y, 214, 170, 'tile', 'tileline', lw=1, name=f'AWARE tile {k+1}')
        t.text_frame.margin_left = px(20); t.text_frame.margin_top = px(14)
        write(t, [[(L, dict(size=66, bold=True, color='teal', font=HEAD))], [(w, dict(size=17, bold=True, color='white'))]])

def L_aware(sc):
    s = new(sc, dark=True)
    text(s, 75, 60, 600, 24, 'Your mental model', name='Label', size=14, color='green')
    title_dark(s, sc['title'], y=88)
    aware_tiles(s, sc['letters'])
    notes(s, *std_notes(sc, 'Each tile fades in when the narrator says its word.'))

def L_opener(sc):
    s = new(sc, dark=True)
    text(s, 75, 70, 400, 150, f"{sc['sectionNo']:02d}", name='Section number', size=110, bold=True, color='teal', font=HEAD)
    title_dark(s, sc['title'], y=230, size=44)
    for k, t in enumerate(sc['lines']):
        text(s, 75, 310 + k * 44, 1000, 40, [('—  ', dict(color='green', bold=True)), (t, {})], name=f'Line {k+1}', size=20, color='white')
    if sc.get('atlasLine'):
        text(s, 75, 310 + len(sc['lines']) * 44 + 10, 1000, 40, sc['atlasLine'], name='Atlas line', size=17, color='lightblue')
    b = box(s, 75, 560, 240, 38, 'teal', radius=0.25, name='Meet Meera button')
    write(b, "Meet Meera’s moment", size=13, bold=True, color='navy', align='c', anchor='m')
    notes(s, *std_notes(sc, 'Base layer. Show layer "Meera" when the timeline ends; the button reopens it.'))
    m = sc['meera']
    ly = layer_slide(sc, 'Meera', f"MEERA VOICEOVER\n{m['vo']}\n\nTriggers: each choice shows its reply and the Continue button; Continue hides this layer; Next unlocks when any choice is Selected.", dark=True)
    ov = box(ly, 0, 36, 1280, 684, 'navy', shape=MSO_SHAPE.RECTANGLE, name='Dim overlay', alpha=70)
    box(ly, 120, 110, 1040, 470, 'white', radius=0.04, name='Meera panel')
    icon(ly, 'meera', 150, 150, 170, 'Meera')
    text(ly, 150, 330, 170, 24, 'Meera', name='Meera name', size=18, bold=True, color='navy', font=HEAD, align='c')
    text(ly, 150, 356, 170, 20, 'Talent Partner', name='Meera role', size=11, color='grey', align='c')
    text(ly, 360, 140, 770, 70, m['situation'], name='Situation', size=16, color='navy')
    text(ly, 360, 215, 770, 30, m['question'], name='Question', size=19, bold=True, color='navy', font=HEAD)
    for k, o in enumerate(m['options']):
        c = box(ly, 360, 255 + k * 50, 770, 42, 'white', 'line', name=f'Choice {k+1}')
        c.text_frame.margin_left = px(16)
        write(c, o['label'], size=15, color='navy', anchor='m')
    for k, o in enumerate(m['options']):
        r = box(ly, 360, 412, 770, 64, 'bluesoft', name=f'Reply {k+1} (Hidden)')
        r.text_frame.margin_left = r.text_frame.margin_right = px(16)
        write(r, o['reply'], size=14, color='navy', anchor='m')
    cb = box(ly, 360, 490, 130, 36, 'teal', radius=0.2, name='Continue (Hidden)')
    write(cb, 'Continue', size=13, bold=True, color='navy', align='c', anchor='m')

def L_defineCards(sc):
    s = new(sc)
    pos = [(75, 150), (739, 150), (75, 330), (739, 330)]
    for k, ((x, y), t) in enumerate(zip(pos, sc['caps'])):
        b = box(s, x, y, 466, 150, 'bluesoft', 'blue', dash=True, name=f'Capability {k+1}')
        b.text_frame.margin_left = b.text_frame.margin_right = px(24)
        write(b, t, size=20, bold=True, color='navy', align='l' if k % 2 == 0 else 'r', anchor='m')
    c = box(s, 455, 220, 370, 200, 'navy', radius=0.1, name='Definition card')
    c.text_frame.margin_left = c.text_frame.margin_right = px(18)
    write(c, [[('AI', dict(size=12, bold=True, color='teal', font=MONO))], [(sc['definition'], dict(size=17, color='white'))]], align='c', anchor='m', spacing=4)
    notes(s, *std_notes(sc))

def L_umbrella(sc):
    s = new(sc)
    u = box(s, 75, 150, 1130, 100, 'navy', shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE, name='AI umbrella')
    write(u, 'AI', size=40, bold=True, color='teal', font=HEAD, align='c', anchor='m')
    for k in range(sc['tiles']):
        b = box(s, 75 + k * 192, 270, 178, 90, 'bluesoft', 'blue', dash=True, name=f'System tile {k+1}')
        icon(s, 'sys', 75 + k * 192 + 71, 297, 36, f'System icon {k+1}')
    text(s, 75, 390, 1100, 60, [(sc['statement'] + ' ', dict(bold=True)), (sc['sub'], {})], name='Statement', size=22, color='navy')
    notes(s, *std_notes(sc))

def L_radial(sc):
    s = new(sc)
    text(s, 75, 125, 1100, 30, sc['definition'], name='Definition', size=20, color='navy')
    h = box(s, 555, 280, 170, 170, 'navy', shape=MSO_SHAPE.OVAL, name='Gen AI hub')
    write(h, 'Gen AI', size=22, bold=True, color='teal', font=HEAD, align='c', anchor='m')
    pos = [(300, 220), (830, 220), (230, 420), (900, 420), (560, 500)]
    for (x, y), t in zip(pos, sc['types']):
        b = box(s, x, y, 150, 46, 'white', 'blue', dash=True, radius=0.5, name=f'Output {t}')
        write(b, t, size=14, bold=True, color='blue', font=MONO, align='c', anchor='m')
    notes(s, *std_notes(sc))

def L_flow(sc):
    s = new(sc)
    text(s, 75, 125, 1130, 60, sc['text'], name='Statement', size=20, color='navy')
    xs = [75, 475, 875]
    d = box(s, xs[0], 260, 330, 150, 'white', 'line', name='Data')
    write(d, [[('• • • • • •', dict(size=16, color='blue'))], [('Data', dict(size=20, bold=True, font=HEAD))]], color='navy', align='c', anchor='m')
    m = box(s, xs[1], 260, 330, 150, 'navy', name='Language model')
    write(m, 'Language model', size=20, bold=True, color='white', font=HEAD, align='c', anchor='m')
    r = box(s, xs[2], 260, 330, 150, 'bluesoft', 'blue', dash=True, name='Response')
    write(r, [[('——————', dict(size=14, color='lightblue'))], [('Response', dict(size=20, bold=True, font=HEAD))]], color='navy', align='c', anchor='m')
    arrow(s, 415, 318, 'blue'); arrow(s, 815, 318, 'blue')
    notes(s, *std_notes(sc))

def L_reveal(sc):
    s = new(sc)
    n = len(sc['cards']); w = (1130 - (n - 1) * 22) / n
    icons = sc.get('icons') or []
    for k, (t, d) in enumerate(sc['cards']):
        x = 75 + k * (w + 22)
        b = box(s, x, 150, w, 230, 'white', 'line', name=f'Card {k+1} (state Normal: term only)')
        b.text_frame.margin_left = px(24); b.text_frame.margin_top = px(22)
        write(b, [[(t, dict(size=18, bold=True, color='blue', font=MONO))], [('Select to reveal', dict(size=11, color='grey'))]], spacing=8)
        o = box(s, x, 400, w, 150, 'bluesoft', 'blue', dash=True, name=f'Card {k+1} OPEN state (copy into the Open state, then delete)')
        o.text_frame.margin_left = px(24); o.text_frame.margin_top = px(16)
        write(o, [[(t, dict(size=16, bold=True, color='blue', font=MONO))], [(d, dict(size=17, color='navy'))]], spacing=6)
    text(s, 75, 565, 600, 22, sc['instruction'], name='Instruction', size=13, color='grey')
    notes(s, *std_notes(sc, 'Top row = cards as learners first see them. Bottom row = how each card looks in its "Open" state: select a card, Edit States > New State "Open", paste the bottom version in, Done. Then delete the bottom row.'))

def q_layout(s, sc, tag):
    text(s, 75, 125, 430, 20, tag, name='Tag', size=11, bold=True, color='grey')
    text(s, 75, 160, 400, 300, [[(l, {})] for l in sc['q']['stem']], name='Question', size=20, bold=True, color='navy', font=HEAD, spacing=12)
    text(s, 75, 470, 400, 20, 'Select one answer.', name='Hint', size=13, color='grey')
    for k, o in enumerate(sc['q']['opts']):
        y = 150 + k * 82
        r = box(s, 515, y, 690, 70, 'white', 'line', name=f'Option {"ABCD"[k]}')
        r.text_frame.margin_left = px(70); r.text_frame.margin_right = px(14)
        write(r, o, size=15, color='navy', anchor='m')
        l = box(s, 530, y + 17, 36, 36, 'white', 'line', radius=0.2, name=f'Option {"ABCD"[k]} letter')
        write(l, 'ABCD'[k], size=14, bold=True, color='navy', font=MONO, align='c', anchor='m')

def fb_layer(sc, name, ok, body, extra=None, how=''):
    ly = layer_slide(sc, name, how)
    box(ly, 0, 36, 1280, 684, 'navy', shape=MSO_SHAPE.RECTANGLE, name='Dim overlay', alpha=45)
    p = box(ly, 260, 180, 760, 360, 'white', radius=0.05, name=f'{name} panel')
    band = box(ly, 260, 180, 760, 10, 'green' if ok else 'coral', shape=MSO_SHAPE.RECTANGLE, name=f'{name} band')
    text(ly, 300, 212, 680, 44, {'Correct': 'Correct', 'Try Again': 'Not Quite', 'Incorrect': 'Not Quite'}.get(name, name),
         name=f'{name} heading', size=28, bold=True, color='greentxt' if ok else 'coralink', font=HEAD)
    text(ly, 300, 265, 680, 100, body, name=f'{name} text', size=17, color='navy')
    if extra: text(ly, 300, 370, 680, 90, extra, name='Answer line', size=14, bold=True, color='navy')
    b = box(ly, 300, 478, 160, 38, 'navy', radius=0.2, name=f'{"Try Again" if name == "Try Again" else "Continue"} button')
    write(b, 'Try Again' if name == 'Try Again' else 'Continue', size=13, bold=True, color='white', align='c', anchor='m')
    return ly

def L_kc(sc):
    s = new(sc)
    q_layout(s, sc, 'Knowledge check · not scored')
    notes(s, *std_notes(sc, 'Quickest: Insert > Convert to Freeform > Pick One and pick the four Option shapes (tick ' + sc['key'] + ' as correct). Attempts 2. Leave it off the results slide.'))
    ans = f"The answer is {sc['key']}: {sc['q']['opts']['ABCD'.index(sc['key'])]}."
    fb_layer(sc, 'Correct', True, sc['fb']['c'], how=f"VOICEOVER\n{sc['fb']['c']}")
    fb_layer(sc, 'Try Again', False, sc['fb']['w'], how=f"VOICEOVER\n{sc['fb']['w']}")
    fb_layer(sc, 'Incorrect', False, sc['fb']['w'], extra=ans, how=f"VOICEOVER\n{sc['fb']['w']}")

def L_quiz(sc):
    s = new(sc)
    q_layout(s, sc, f"Question {sc['qNo']} of 8 · Attempt %QuizAttempts% of 3")
    var = ['Missed_S1', 'Missed_S2', 'Missed_S4', 'Missed_S5', 'Missed_S6', 'Missed_S4', 'Missed_S3', 'Missed_S7'][sc['qNo'] - 1]
    notes(s, *std_notes(sc, f"Graded question: tick {sc['key']} as correct, 10 points, Attempts 1, Shuffle Answers. Incorrect layer trigger: set {var} = True when its timeline starts."))
    fb_layer(sc, 'Correct', True, sc['fb']['c'], how=f"VOICEOVER\n{sc['fb']['c']}")
    fb_layer(sc, 'Incorrect', False, sc['fb']['w'], how=f"VOICEOVER\n{sc['fb']['w']}\n\nTrigger: set {var} = True when the timeline starts on this layer.")

def L_beforeAfter(sc):
    s = new(sc)
    for k, (t, d) in enumerate(sc['cards']):
        x = 75 + k * 580
        box(s, x + 120, 140, 100, 132, 'white', 'line', shape=MSO_SHAPE.RECTANGLE, name=f'Page before {k+1}')
        arrow(s, x + 240, 190, 'blue')
        pg = box(s, x + 300, 140, 100, 132, 'bluesoft', 'blue', dash=True, shape=MSO_SHAPE.RECTANGLE, name=f'Page after {k+1}')
        write(pg, '———\n———\n——', size=12, color='blue', align='c', anchor='m')
        card(s, x, 300, 550, 120, t, d)
    notes(s, *std_notes(sc))

def L_condense(sc):
    s = new(sc)
    box(s, 330, 130, 100, 180, 'white', 'line', shape=MSO_SHAPE.RECTANGLE, name='Long document')
    arrow(s, 450, 205, 'blue')
    box(s, 510, 175, 100, 90, 'bluesoft', 'blue', dash=True, shape=MSO_SHAPE.RECTANGLE, name='Summary page')
    arrow(s, 630, 205, 'blue')
    l = box(s, 690, 140, 100, 160, 'bluesoft', 'blue', dash=True, shape=MSO_SHAPE.RECTANGLE, name='List page')
    write(l, '• ——\n• ——\n• ——', size=12, color='blue', align='c', anchor='m')
    for k, (t, d) in enumerate(sc['cards']):
        card(s, 75 + k * 580, 340, 550, 120, t, d)
    notes(s, *std_notes(sc))

def L_pattern(sc):
    s = new(sc)
    for k, (x, y) in enumerate([(110, 160), (330, 300), (220, 430)]):
        c = box(s, x, y, 130, 130, 'bluesoft', 'blue', dash=True, shape=MSO_SHAPE.OVAL, name=f'Cluster {k+1}')
        write(c, '● ●\n●', size=16, color='blue', align='c', anchor='m')
    card(s, 620, 270, 585, 150, sc['card'][0], sc['card'][1], term_size=16, desc_size=20)
    notes(s, *std_notes(sc, 'Optional: start the dots scattered and give them short motion paths into the clusters.'))

def L_limits(sc):
    s = new(sc)
    text(s, 75, 150, 330, 50, sc['intro'], name='Intro', size=28, bold=True, color='navy', font=HEAD)
    for k, t in enumerate(sc['items']):
        y = 145 + k * 78
        x = box(s, 440, y + 10, 36, 36, 'white', 'coral', shape=MSO_SHAPE.OVAL, name=f'Cross {k+1}')
        write(x, '✕', size=13, bold=True, color='coralink', align='c', anchor='m')
        text(s, 495, y + 8, 700, 40, t, name=f'Limit {k+1}', size=24, bold=True, color='navy', font=HEAD, anchor='m')
        line(s, 440, y + 62, 1205, y + 62, 'line', 1, f'Limit line {k+1}')
    notes(s, *std_notes(sc))

def L_notEqual(sc):
    s = new(sc)
    l = box(s, 75, 230, 440, 150, 'bluesoft', 'blue', dash=True, name='Confident output')
    write(l, sc['left'], size=24, bold=True, color='blue', font=MONO, align='c', anchor='m')
    r = box(s, 765, 230, 440, 150, 'greensoft', 'greentxt', name='Correct output')
    write(r, sc['right'], size=24, bold=True, color='greentxt', font=MONO, align='c', anchor='m')
    text(s, 560, 225, 160, 160, '≠', name='Not equal symbol', size=110, bold=True, color='coral', align='c', anchor='m')
    notes(s, *std_notes(sc, 'The ≠ fades in last. Right-click it > Save as Picture to reuse on screen 26.'))

def L_pair(sc):
    s = new(sc)
    sym = {'idea': '☀', 'draft': '✎', 'doc': '☰', 'search': '⌕'}
    for k, (t, d) in enumerate(sc['cards']):
        b = card(s, 75 + k * 580, 150, 550, 260, t, d, term_size=16, desc_size=22)
        b.text_frame.margin_top = px(80)
        icon(s, sc['icons'][k] if sc['icons'][k] in sym else 'doc', 100 + k * 580, 172, 44, f'Icon {k+1}')
    notes(s, *std_notes(sc))

def L_twoCol(sc):
    s = new(sc)
    card(s, 75, 150, 550, 230, sc['left'][0], sc['left'][1], 'ai', 16, 23)
    card(s, 655, 150, 550, 230, sc['right'][0], sc['right'][1], 'human', 16, 23)
    o = box(s, 200, 420, 150, 50, 'white', 'blue', dash=True, name='Output card (motion path to the right)')
    write(o, '————', size=12, color='lightblue', align='c', anchor='m')
    t = box(s, 900, 423, 44, 44, 'teal', shape=MSO_SHAPE.OVAL, name='Review tick')
    write(t, '✓', size=18, bold=True, color='navy', align='c', anchor='m')
    notes(s, *std_notes(sc, 'Give the output card a motion path (Lines, right) ending under the HUMAN card; the tick fades in after it.'))

def three_steps(s, items, y=140, h=170, lock=False):
    w = 340
    for k, (t, d) in enumerate(items):
        x = 75 + k * (w + 55)
        b = box(s, x, y, w, h, 'white', 'line', name=f'Step {k+1}')
        b.text_frame.margin_left = b.text_frame.margin_right = px(22); b.text_frame.margin_top = px(18)
        head = [(('LOCKED' if lock else f'STEP {k+1}'), dict(size=11, bold=True, color='grey'))]
        if lock: icon(s, 'lockGrey', x + w - 50, y + 16, 28, f'Lock {k+1}')
        write(b, [head, [(t, dict(size=16, bold=True, color='blue', font=MONO))], [(d, dict(size=17, color='navy'))]], spacing=6)
        if k: arrow(s, x - 48, y + h / 2 - 15)

def L_steps(sc):
    s = new(sc)
    three_steps(s, sc['steps'])
    wc = box(s, 75, 360, 260, 110, 'white', 'blue', dash=True, name='Polished answer card')
    write(wc, '———————\n—————\n——————', size=12, color='lightblue', align='c', anchor='m')
    m = box(s, 312, 345, 40, 40, 'coral', shape=MSO_SHAPE.OVAL, name='Not-equal badge (Pulse once)')
    write(m, '≠', size=18, bold=True, color='white', align='c', anchor='m')
    text(s, 380, 380, 820, 70, sc['warning'], name='Warning', size=22, bold=True, color='navy', font=HEAD, anchor='m')
    notes(s, *std_notes(sc))

def L_hotspot(sc):
    s = new(sc)
    text(s, 75, 140, 900, 26, sc['intro'], name='Intro', size=18, color='navy')
    b = box(s, 75, 180, 1000, 150, 'bluesoft', 'blue', dash=True, name='AI draft')
    b.text_frame.margin_left = b.text_frame.margin_right = px(24); b.text_frame.margin_top = px(16)
    write(b, [[('AI DRAFT', dict(size=12, bold=True, color='blue', font=MONO))], [(sc['text'], dict(size=19, color='navy', font=MONO))]], spacing=8)
    text(s, 75, 350, 900, 26, sc['prompt'], name='Prompt', size=18, bold=True, color='navy')
    sb = box(s, 75, 395, 130, 38, 'navy', radius=0.2, name='Submit button')
    write(sb, 'Submit', size=13, bold=True, color='white', align='c', anchor='m')
    for k, c in enumerate(sc['claims']):
        cl = box(s, 760, 440 + k * 46, 300, 38, 'tealsoft', 'teal', name=f'Claim {k+1} box: draw over “{c}”')
        write(cl, f'Claim {k+1}: {c}', size=12, color='tealink', align='c', anchor='m')
    text(s, 760, 410, 440, 24, 'Move each claim box over its words in the draft, then set its fill to 100% transparent.', name='DELETE ME: claim note', size=10, italic=True, color='grey')
    notes(s, *std_notes(sc, 'Claim boxes: Hover and Selected states. Submit: add 1 to Tries30a; Correct if all three Selected; else Incorrect. Incorrect: if Tries30a >= 2 select all claims, hide Try Again, show Continue.'))
    fb_layer(sc, 'Correct', True, sc['pfb']['c'], how=f"VOICEOVER\n{sc['pfb']['c']}")
    fb_layer(sc, 'Incorrect', False, sc['pfb']['w'], how=f"VOICEOVER\n{sc['pfb']['w']}")

def L_balance(sc):
    s = new(sc)
    icon(s, 'balance', 105, 220, 160, 'Balance icon')
    card(s, 330, 170, 425, 230, sc['cards'][0][0], sc['cards'][0][1], 'ai', 16, 21)
    card(s, 780, 170, 425, 230, sc['cards'][1][0], sc['cards'][1][1], 'human', 16, 21)
    notes(s, *std_notes(sc))

def L_checks3(sc):
    s = new(sc)
    text(s, 75, 140, 1000, 30, sc['intro'], name='Intro', size=20, color='navy')
    for k, (t, d) in enumerate(sc['checks']):
        x = 75 + k * 385
        b = box(s, x, 190, 360, 230, 'white', 'line', name=f'Check {k+1}')
        b.text_frame.margin_left = b.text_frame.margin_right = px(24); b.text_frame.margin_top = px(84)
        write(b, [[(t, dict(size=18, bold=True, color='blue', font=MONO))], [(d, dict(size=17, color='navy'))]], spacing=6)
        c = box(s, x + 24, 210, 52, 52, 'greentxt', shape=MSO_SHAPE.OVAL, name=f'Tick {k+1}')
        write(c, '✓', size=20, bold=True, color='white', align='c', anchor='m')
    notes(s, *std_notes(sc, '⚑ Pending Legal review: build it, hold the voice.'))

def L_gate(sc):
    s = new(sc)
    line(s, 75, 210, 330, 210, 'blue', 3, 'AI path')
    text(s, 75, 180, 60, 22, 'AI', name='AI label', size=12, bold=True, color='blue', font=MONO)
    box(s, 332, 180, 14, 60, 'coral', shape=MSO_SHAPE.RECTANGLE, name='Gate')
    line(s, 75, 320, 470, 320, 'teal', 5, 'Human path')
    text(s, 75, 290, 100, 22, 'HUMAN', name='Human label', size=12, bold=True, color='tealink', font=MONO)
    arrow(s, 460, 302)
    text(s, 540, 170, 665, 90, sc['text'], name='Statement', size=24, bold=True, color='navy', font=HEAD)
    text(s, 540, 270, 665, 70, sc['sub'], name='Sub', size=19, color='navy')
    b = box(s, 75, 470, 1130, 80, 'navy', radius=0.12, name='Principle band')
    b.text_frame.margin_left = px(30)
    write(b, [(sc['closing'][0] + ' ', dict(color='lightblue')), (sc['closing'][1], dict(color='teal'))], size=30, bold=True, font=HEAD, anchor='m')
    notes(s, *std_notes(sc, '⚑ Pending Legal/HR review of the voiceover. Copy the band for screen 66.'))

def L_checklist(sc):
    s = new(sc)
    for k, t in enumerate(sc['items']):
        y = 145 + k * 84
        r = box(s, 75, y, 900, 70, 'white', 'line', name=f'Question {k+1}')
        r.text_frame.margin_left = px(84)
        write(r, [(f'0{k+1}  ', dict(size=12, bold=True, color='grey', font=MONO)), (t, dict(size=22, bold=True, font=HEAD))], color='navy', anchor='m')
        box(s, 95, y + 13, 44, 44, 'white', 'line', radius=0.2, name=f'Tick box {k+1}')
    text(s, 75, 490, 600, 22, sc['instruction'], name='Instruction', size=13, color='grey')
    tk = box(s, 1000, 145, 200, 70, 'greensoft', 'greentxt', name='TICKED state example (copy into each row’s "Ticked" state, then delete)')
    write(tk, '✓ Ticked state', size=14, bold=True, color='greentxt', align='c', anchor='m')
    notes(s, *std_notes(sc))

def L_categories(sc):
    s = new(sc)
    for k, t in enumerate(sc['items']):
        x = 75 + (k % 4) * 288; y = 160 + (k // 4) * 150
        b = box(s, x, y, 266, 125, 'navy', name=f'Category {t}')
        b.text_frame.margin_top = px(46)
        write(b, t, size=15, bold=True, color='white', font=MONO, align='c', anchor='m')
        icon(s, 'lock', x + 115, y + 24, 36, f'Lock {t}')
    notes(s, *std_notes(sc, '⚑ Pending Privacy check of Client, Proprietary, Credentials.'))

def L_threeGate(sc):
    s = new(sc)
    for k, t in enumerate(sc['parts']):
        x = 75 + k * 400
        b = box(s, x, 200, 340, 170, 'white', 'greentxt', name=f'Gate {k+1}')
        b.text_frame.margin_top = px(50)
        write(b, t, size=18, bold=True, color='navy', font=MONO, align='c', anchor='m')
        icon(s, 'open', x + 150, 228, 40, f'Open lock {k+1}')
        if k: text(s, x - 52, 255, 40, 50, '+', name='Plus', size=34, bold=True, color='grey', align='c')
    notes(s, *std_notes(sc))

def L_stages(sc):
    s = new(sc)
    for k, t in enumerate(sc['stages']):
        cx = 220 + k * 420
        c = box(s, cx - 48, 180, 96, 96, 'navy', shape=MSO_SHAPE.OVAL, name=f'Lock {k+1}')
        icon(s, 'lock', cx - 22, 206, 44, f'Lock icon {k+1}')
        text(s, cx - 150, 290, 300, 30, t, name=f'Stage {k+1}', size=20, bold=True, color='navy', font=HEAD, align='c')
        if k: line(s, cx - 372, 228, cx - 60, 228, 'teal', 3, f'Stage line {k}')
    text(s, 75, 380, 1130, 30, sc['sub'], name='Sub', size=20, bold=True, color='navy', align='c')
    notes(s, *std_notes(sc))

def L_recover(sc):
    s = new(sc)
    three_steps(s, sc['steps'], lock=True)
    text(s, 75, 330, 900, 24, 'Replace [approved contact] with the real contact before recording.', name='DELETE ME: pending note', size=12, italic=True, color='coralink')
    notes(s, *std_notes(sc, '⚑ Pending: the real contact and process. Hold the voice.'))

def L_promptPart(sc):
    s = new(sc, eyebrow=f"Prompt element {sc['part']+1} of 5")
    t, d = [x.strip() for x in sc['onscreen'].split('\n')][:2]
    c = box(s, 230, 160, 820, 200, 'white', 'line', name='Element card')
    write(c, [[(t, dict(size=24, bold=True, color='blue', font=MONO))], [(d, dict(size=28, bold=True, color='navy', font=HEAD))]], align='c', anchor='m', spacing=8)
    text(s, 75, 470, 400, 18, 'Your prompt so far', name='Builder label', size=11, color='grey')
    for k, p in enumerate(sc['parts']):
        on = k <= sc['part']
        b = box(s, 75 + k * 228, 495, 218, 56, 'blue' if on else None, 'blue' if on else 'line', dash=not on, name=f'Slot {p} ({"Filled" if on else "Empty"})')
        write(b, p, size=13, bold=True, color='white' if on else 'grey', font=MONO, align='c', anchor='m')
    notes(s, *std_notes(sc, 'Slots have two states: Empty (dashed grey) and Filled (Pro Blue).'))

def L_assemble(sc):
    s = new(sc)
    for k, p in enumerate(sc['parts']):
        b = box(s, 75 + k * 226, 220, 226, 80, 'blue', shape=MSO_SHAPE.RECTANGLE, name=f'Slot {p}')
        write(b, p, size=14, bold=True, color='white', font=MONO, align='c', anchor='m')
    text(s, 75, 340, 1130, 80, sc['sub'], name='Sub', size=26, bold=True, color='navy', font=HEAD, align='c')
    notes(s, *std_notes(sc))

def L_compare(sc):
    s = new(sc)
    card(s, 75, 150, 400, 260, 'VAGUE', sc['vague'], 'white', 14, 22)
    card(s, 500, 150, 705, 260, 'CLEARER', sc['clear'], 'ai', 14, 22)
    text(s, 75, 430, 1130, 60, 'Highlights to add on cue: ' + ' · '.join(f'“{w}” = {l}' for w, l in sc['highlights']),
         name='DELETE ME: highlight list', size=12, italic=True, color='grey')
    notes(s, *std_notes(sc, 'For each added detail draw a light-teal rectangle behind the words with a tiny label; fade each on its cue.'))

def L_dragPrompt(sc):
    s = new(sc)
    b = box(s, 75, 140, 540, 64, 'white', 'line', name='Start with')
    b.text_frame.margin_left = px(16)
    write(b, [[('START WITH', dict(size=11, bold=True, color='blue', font=MONO))], [(sc['base'], dict(size=17, color='navy', font=MONO))]], anchor='m')
    for k, sl in enumerate(sc['slots']):
        d = box(s, 75, 222 + k * 82, 540, 70, None, 'line', dash=True, name=f'Drop slot {sl["part"]}')
        d.text_frame.margin_left = px(16)
        write(d, [[(sl['part'], dict(size=12, bold=True, color='blue', font=MONO))], [('Drop a piece here', dict(size=12, color='grey'))]], anchor='m')
    text(s, 660, 140, 540, 20, 'Drag the missing pieces into the prompt', name='Bank label', size=13, bold=True, color='grey')
    pieces = []
    for sl in sc['slots']: pieces += [(sl['right'], True), (sl['wrong'], False)]
    order = [0, 3, 5, 6, 1, 2, 4, 7]
    for k, i in enumerate(order):
        t, ok = pieces[i]
        x = 660 + (k % 2) * 275; y = 172 + (k // 2) * 78
        p = box(s, x, y, 262, 66, 'white', 'line', name=f'Piece {k+1}{" (correct)" if ok else " (distractor)"}')
        p.text_frame.margin_left = p.text_frame.margin_right = px(10)
        write(p, t, size=12, color='navy', font=MONO, anchor='m')
    notes(s, *std_notes(sc, 'Insert > Convert to Freeform > Drag and Drop. Pair each correct piece with its slot; distractors get no target. Attempts 2.'))
    strong = f"{sc['slots'][0]['right']} {sc['base']} " + ' '.join(x['right'] for x in sc['slots'][1:])
    fb_layer(sc, 'Correct', True, sc['pfb']['c'], extra='Stronger prompt: ' + strong, how=f"VOICEOVER\n{sc['pfb']['c']}")
    fb_layer(sc, 'Incorrect', False, sc['pfb']['w'], extra='Stronger prompt: ' + strong, how=f"VOICEOVER\n{sc['pfb']['w']}")

def L_agentModel(sc):
    s = new(sc)
    text(s, 75, 130, 1130, 30, sc['text'], name='Statement', size=20, color='navy')
    for k, t in enumerate(sc['inputs']):
        b = box(s, 75, 200 + k * 85, 480, 66, 'bluesoft', 'blue', dash=True, name=f'Input {k+1}')
        b.text_frame.margin_left = px(20)
        write(b, t, size=16, bold=True, color='blue', font=MONO, anchor='m')
        line(s, 560, 233 + k * 85, 760, 318, 'blue', 2, f'Connector {k+1}')
    t = box(s, 780, 255, 425, 130, 'navy', name='Defined task')
    write(t, 'A defined task', size=26, bold=True, color='white', font=HEAD, align='c', anchor='m')
    notes(s, *std_notes(sc))

def L_lanes(sc):
    s = new(sc)
    for k, (who, pills, cap) in enumerate([('CHATBOT', ['You ask', 'AI responds'], sc['chatbot']), ('AGENT', ['You define a task', 'Step 1', 'Step 2', 'Step 3'], sc['agent'])]):
        y = 150 + k * 190
        box(s, 75, y, 1130, 160, 'white', 'line', name=f'{who} lane')
        text(s, 100, y + 50, 150, 30, who, name=f'{who} label', size=16, bold=True, color='blue', font=MONO)
        x = 260
        for j, p in enumerate(pills):
            w = 70 + len(p) * 10
            b = box(s, x, y + 40, w, 46, 'blue' if j else 'bluesoft', 'blue', dash=not j, radius=0.5, name=f'{who} pill {j+1}')
            write(b, p, size=14, bold=True, color='white' if j else 'navy', align='c', anchor='m')
            x += w + 14
            if j < len(pills) - 1: arrow(s, x - 12, y + 48, 'grey', 18); x += 24
        text(s, 260, y + 100, 900, 26, cap, name=f'{who} caption', size=14, color='grey')
    notes(s, *std_notes(sc))

def L_atlasVideo(sc):
    s = new(sc)
    x = 75
    for k, p in enumerate(sc['process']):
        b = box(s, x, 130, 150, 34, 'teal' if k == 3 else 'navy', radius=0.2, name=f'Process {p}')
        write(b, p, size=12, bold=True, color='navy' if k == 3 else 'white', font=MONO, align='c', anchor='m')
        x += 150
        if k < 3: arrow(s, x + 2, 132, 'grey', 16); x += 44
    v = box(s, 75, 185, 1130, 400, 'tile', 'tileline', dash=True, name='VIDEO PLACEHOLDER (replace with Insert > Video)')
    write(v, [[('▶', dict(size=40, color='teal'))], [('Your Atlas screen recording goes here', dict(size=18, color='white'))], [(sc['note'], dict(size=13, color='lightblue'))]], align='c', anchor='m', spacing=6)
    notes(s, *std_notes(sc, 'Insert > Video > Video from File. Next unlocks when media completes.'))

def L_atlasSteps(sc):
    s = new(sc)
    for k, t in enumerate(sc['steps']):
        x = 75 + k * 290
        b = box(s, x, 160, 270, 260, 'tealsoft' if k == 3 else 'white', 'teal' if k == 3 else 'line', name=f'Step {k+1}')
        b.text_frame.margin_left = b.text_frame.margin_right = px(20); b.text_frame.margin_top = px(20)
        write(b, [[(f'STEP {k+1}', dict(size=11, bold=True, color='grey', font=MONO))], [(t, dict(size=22, bold=True, color='navy', font=HEAD))], [(f'▶ Screen recording {k+1}', dict(size=11, color='grey'))]], spacing=14)
    notes(s, *std_notes(sc))

def L_tryIt(sc):
    s = new(sc)
    for k, (t, d) in enumerate(sc['tabs']):
        b = box(s, 75, 140 + k * 92, 290, 80, 'white', 'line', name=f'Tab {t}')
        b.text_frame.margin_left = px(14)
        write(b, [[(t, dict(size=12, bold=True, color='blue', font=MONO))], [(d, dict(size=13, color='navy'))]], anchor='m')
    text(s, 75, 515, 290, 20, 'Open all four tabs, then try it in Atlas.', name='Instruction', size=11, color='grey')
    p = box(s, 390, 140, 815, 440, 'white', 'line', name='Tab panel')
    p.text_frame.margin_left = px(26); p.text_frame.margin_top = px(24)
    write(p, [[('Select a Tab to Begin', dict(size=22, bold=True, font=HEAD))], [('Each tab shows one part of the activity: the scenario, the sample notes, the prompt to use in Atlas, and how to review the output.', dict(size=15))]], color='navy', spacing=8)
    notes(s, *std_notes(sc, 'Four layers (Scenario, Input, Task, Output) with "Hide other slide layers" ticked. Next unlocks when all four tabs are Visited.'))
    def panel(ly, paras):
        q = box(ly, 390, 140, 815, 440, 'white', 'line', name='Tab panel content')
        q.text_frame.margin_left = q.text_frame.margin_right = px(26); q.text_frame.margin_top = px(22)
        write(q, paras, color='navy', spacing=6)
    ly = layer_slide(sc, 'Scenario', 'Shown when the Scenario tab is clicked.')
    panel(ly, [[(sc['tabs'][0][1].title() if False else 'Meeting Action Tracker', dict(size=22, bold=True, font=HEAD))], [(sc['scenario'], dict(size=16))], [(sc['tabs'][3][1] + '.', dict(size=16, bold=True))]])
    ly = layer_slide(sc, 'Input', 'Shown when the Input tab is clicked.')
    panel(ly, [[('Sample Meeting Notes', dict(size=22, bold=True, font=HEAD))]] + [[('• ' + t, dict(size=13.5, bold=(t == sc['noOwnerLine'])))] for t in sc['notes']] + [[('Names have been changed.', dict(size=11, color='grey'))]])
    ly = layer_slide(sc, 'Task', 'Shown when the Task tab is clicked.')
    panel(ly, [[('Prompt for Atlas', dict(size=22, bold=True, font=HEAD))], [(sc['prompt'], dict(size=13, font=MONO))]])
    ly = layer_slide(sc, 'Output', 'Shown when the Output tab is clicked.')
    rows = [[(f"{a}  |  {o}  |  {d}", dict(size=12.5, font=MONO))] for a, o, d in sc['expected']]
    panel(ly, [[('Expected Action Items', dict(size=22, bold=True, font=HEAD))], [('Action  |  Owner  |  Due date', dict(size=12, bold=True, font=MONO))]] + rows +
          [[('Before you use the output, ask:', dict(size=14, bold=True))]] + [[(f'{i+1}. {q}', dict(size=13))] for i, q in enumerate(sc['check'])])

def L_assessIntro(sc):
    s = new(sc, dark=True)
    text(s, 75, 150, 600, 24, 'Comprehensive Assessment', name='Label', size=14, color='green')
    title_dark(s, 'Show What You Know.', y=180, size=60)
    text(s, 75, 290, 1000, 40, sc['lines'][1], name='Sub', size=22, color='white')
    x = 75
    for t in sc['info']:
        w = 40 + len(t) * 10
        b = box(s, x, 370, w, 44, None, 'lightblue', radius=0.5, name=f'Rule {t}')
        write(b, t, size=15, bold=True, color='white', align='c', anchor='m'); x += w + 14
    notes(s, *std_notes(sc))

def L_results(sc):
    s = new(sc)
    text(s, 75, 150, 1000, 70, '%Results.ScorePercent%%', name='Score (built-in variable)', size=40, bold=True, color='greentxt', font=HEAD)
    text(s, 75, 250, 600, 24, 'Your score: %Results.ScorePoints% · Passing score: 80%', name='Score line', size=14, color='grey')
    notes(s, *std_notes(sc, 'Graded Result Slide: selected questions = the 8 assessment questions, pass 80%. Success and Failure layers follow.'))
    lines = sc['results'].split('\n')
    passline = [l for l in lines if l.startswith('PASS')][0].replace('PASS', '', 1).strip()
    notyet = [l for l in lines if l.startswith('NOT YET')][0].replace('NOT YET', '', 1).strip()
    ly = layer_slide(sc, 'Success', 'Success layer: Continue jumps to the next slide.')
    text(ly, 520, 160, 680, 50, 'Pass', name='Pass heading', size=36, bold=True, color='navy', font=HEAD)
    text(ly, 520, 220, 680, 40, passline, name='Pass text', size=20, color='navy')
    b = box(ly, 520, 280, 140, 38, 'navy', radius=0.2, name='Continue button'); write(b, 'Continue', size=13, bold=True, color='white', align='c', anchor='m')
    ly = layer_slide(sc, 'Failure', 'Failure layer: Review buttons start Hidden and show if their Missed_S variable is True; Retry hides and the message shows when QuizAttempts >= 3.')
    text(ly, 520, 150, 680, 50, 'Not Yet', name='Not yet heading', size=36, bold=True, color='navy', font=HEAD)
    text(ly, 520, 205, 680, 40, notyet, name='Not yet text', size=18, color='navy')
    labels = ['Understand AI', 'Capabilities and limits', 'AI at work', 'Use AI responsibly', 'Protect information', 'Prompt with purpose', 'AI agents and Atlas']
    for k, l in enumerate(labels):
        x = 520 + (k % 2) * 340; y = 255 + (k // 2) * 46
        r = box(ly, x, y, 320, 38, 'coralsoft', 'coral', radius=0.5, name=f'Review S{k+1} (Hidden)')
        write(r, 'Review: ' + l, size=12, bold=True, color='navy', align='c', anchor='m')
    rb = box(ly, 520, 450, 200, 40, 'navy', radius=0.2, name='Retry assessment'); write(rb, 'Retry assessment', size=13, bold=True, color='white', align='c', anchor='m')
    m = box(ly, 520, 505, 680, 70, 'coralsoft', radius=0.1, name='Attempts used message (Hidden)')
    m.text_frame.margin_left = m.text_frame.margin_right = px(14)
    write(m, sc['outOfAttempts'], size=14, color='navy', anchor='m')

def L_awareClose(sc):
    s = new(sc, dark=True)
    text(s, 75, 70, 600, 24, sc['title'], name='Label', size=16, bold=True, color='green')
    aware_tiles(s, sc['letters'], y=130)
    b = box(s, 75, 360, 1130, 90, 'tile', radius=0.12, name='Principle band')
    b.text_frame.margin_left = px(30)
    write(b, [('AI can assist. ', dict(color='lightblue')), ('Humans remain accountable.', dict(color='teal'))], size=34, bold=True, font=HEAD, anchor='m')
    notes(s, *std_notes(sc))

def L_download(sc):
    s = new(sc)
    c = box(s, 75, 150, 380, 240, 'navy', name='Download card')
    c.text_frame.margin_left = px(28); c.text_frame.margin_top = px(26)
    write(c, [[('DOWNLOAD', dict(size=12, bold=True, color='teal', font=MONO))], [('One-page Level 1 quick reference', dict(size=24, bold=True, color='white', font=HEAD))]], spacing=10)
    b = box(s, 103, 320, 170, 40, 'teal', radius=0.2, name='Download PDF button'); write(b, 'Download PDF', size=13, bold=True, color='navy', align='c', anchor='m')
    for k, (t, d) in enumerate(sc['items']):
        c2 = card(s, 490, 150 + k * 82, 715, 70, t, d, 'white', 11, 15); c2.text_frame.margin_top = px(10)
    notes(s, *std_notes(sc, 'Download button: Jump to URL/file > attach the quick-reference PDF. Also add it in Player > Resources.'))

def L_levels(sc):
    s = new(sc)
    for k, (a, n, d) in enumerate(sc['levels']):
        x = 75 + k * 283
        b = box(s, x, 180, 283, 260, 'navy' if k == 0 else 'white', 'navy' if k == 0 else 'line', shape=MSO_SHAPE.RECTANGLE, name=f'Level {k+1}')
        b.text_frame.margin_left = b.text_frame.margin_right = px(22); b.text_frame.margin_top = px(24)
        write(b, [[(a, dict(size=12, bold=True, color='teal' if k == 0 else 'grey', font=MONO))], [(n, dict(size=22, bold=True, color='white' if k == 0 else 'navy', font=HEAD))], [(d, dict(size=15, color='lightblue' if k == 0 else 'grey'))]], spacing=8)
    t = box(s, 97, 390, 110, 30, 'teal', radius=0.5, name='Complete tag'); write(t, 'Complete', size=11, bold=True, color='navy', align='c', anchor='m')
    notes(s, *std_notes(sc))

for sc in SCREENS:
    globals()['L_' + sc['layout']](sc)

prs.core_properties.title = 'AI Literacy Level 1 – Storyline import'
prs.save(OUT)
print('slides', len(prs.slides))
