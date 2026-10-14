import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY, TEAL, GREEN, BLUE = RGBColor(0x02,0x1F,0x62), RGBColor(0x00,0xBB,0xC2), RGBColor(0x7E,0xBD,0x31), RGBColor(0x49,0x7E,0xDC)
SNOW, GREY, LIGHTBLUE = RGBColor(0xF6,0xF6,0xF9), RGBColor(0x5A,0x5C,0x61), RGBColor(0x9D,0xB8,0xEE)

prs = Presentation(sys.argv[1])
# drop every template slide; the masters and layouts stay
lst = prs.slides._sldIdLst
for sid in list(lst):
    prs.part.drop_rel(sid.rId); lst.remove(sid)

s = prs.slides.add_slide(prs.slide_layouts[5])   # BASIC SLIDE: logo, footer, brand line
s.placeholders[10].text = 'AI Literacy Level 1: Quick Reference'

def text(x, y, w, h, runs, size=14, bold=False, color=NAVY, font=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: tb.name = name
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf, m, 0)
    p = tf.paragraphs[0]; p.alignment = align
    for t, opt in (runs if isinstance(runs, list) else [(runs, {})]):
        r = p.add_run(); r.text = t
        r.font.size = Pt(opt.get('size', size)); r.font.bold = opt.get('bold', bold)
        r.font.color.rgb = opt.get('color', color)
        if opt.get('font', font): r.font.name = opt.get('font', font)
    return tb

def box(x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, name=None):
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name: b.name = name
    b.fill.solid(); b.fill.fore_color.rgb = fill
    if line: b.line.color.rgb = line; b.line.width = Pt(1)
    else: b.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: b.adjustments[0] = 0.12
    b.shadow.inherit = False
    return b

top, colw, gap, x0 = 1.45, 3.62, 0.43, 0.78
cols = [x0 + i * (colw + gap) for i in range(3)]

# column panels
for i, (title, sub) in enumerate([('AWARE', 'Your mental model for working with AI'),
                                  ('Before You Share', 'Ask these four questions before you enter information into an AI tool'),
                                  ('A Clear Prompt', 'Give AI a clearer target with five elements')]):
    box(cols[i], top, colw, 4.35, SNOW, name=f'Panel {i+1}')
    text(cols[i] + 0.25, top + 0.2, colw - 0.5, 0.4, title, size=20, bold=True, color=NAVY, font='Aptos Display', name=f'Heading {i+1}')
    text(cols[i] + 0.25, top + 0.62, colw - 0.5, 0.5, sub, size=11, color=GREY, name=f'Subheading {i+1}')

# AWARE
aware = [('A', 'Appreciate AI'), ('W', 'Weigh the Output'), ('A', 'Assess the Risk'), ('R', 'Respond with the Right Prompt'), ('E', 'Ensure Human Accountability')]
for k, (L, w) in enumerate(aware):
    y = top + 1.22 + k * 0.6
    t = box(cols[0] + 0.25, y, 0.46, 0.46, NAVY, name=f'AWARE tile {k+1}')
    tf = t.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = L
    r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = TEAL; r.font.name = 'Aptos Display'
    text(cols[0] + 0.88, y, colw - 1.1, 0.46, w, size=14, bold=True, anchor=MSO_ANCHOR.MIDDLE, name=f'AWARE word {k+1}')

# Before you share
qs = ['What information is it?', 'Who owns it?', 'Is the tool approved?', 'Am I permitted to use it?']
for k, q in enumerate(qs):
    y = top + 1.32 + k * 0.72
    c = box(cols[1] + 0.25, y, 0.42, 0.42, TEAL, shape=MSO_SHAPE.OVAL, name=f'Number {k+1}')
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = str(k + 1)
    r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = NAVY
    text(cols[1] + 0.85, y, colw - 1.05, 0.42, q, size=14, bold=True, anchor=MSO_ANCHOR.MIDDLE, name=f'Question {k+1}')

# prompt elements
els = [('Context', 'What does AI need to know to interpret the situation?'), ('Task', 'What exactly do you want AI to do?'),
       ('Audience', 'Who is the output for?'), ('Output', 'What should the result look like?'), ('Constraints', 'What boundaries should AI follow?')]
for k, (e, q) in enumerate(els):
    y = top + 1.2 + k * 0.58
    text(cols[2] + 0.25, y, 1.15, 0.42, e, size=12, bold=True, color=BLUE, name=f'Element {k+1}')
    text(cols[2] + 1.42, y, colw - 1.6, 0.55, q, size=11, color=NAVY, name=f'Element question {k+1}')
text(cols[2] + 0.25, top + 4.0, colw - 0.5, 0.3, 'Review the result, then refine.', size=12, bold=True, color=NAVY, name='Refine')

# closing band
band = box(x0, 6.0, 11.77, 0.58, NAVY, name='Principle band')
tf = band.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.margin_left = Inches(0.3)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
for t, c in (('AI can assist. ', LIGHTBLUE), ('Humans remain accountable.', TEAL)):
    r = p.add_run(); r.text = t; r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = c; r.font.name = 'Aptos Display'

s.shapes.title  # no-op; title lives in placeholder 10
prs.core_properties.title = 'AI Literacy Level 1: Quick Reference'
prs.save(sys.argv[2])
print('saved')
